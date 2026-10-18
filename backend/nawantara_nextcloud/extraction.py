"""Bounded local text extraction with English and Indonesian OCR."""

from __future__ import annotations

import io
import subprocess
import tempfile
from pathlib import Path
from zipfile import ZipFile

SUPPORTED = {
    '.txt',
    '.md',
    '.csv',
    '.json',
    '.pdf',
    '.png',
    '.jpg',
    '.jpeg',
    '.tif',
    '.tiff',
    '.docx',
    '.xlsx',
    '.pptx',
}


def ocr(path: Path) -> str:
    return subprocess.run(
        ['tesseract', str(path), 'stdout', '-l', 'eng+ind'], capture_output=True, check=True, timeout=120
    ).stdout.decode()


def pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(path)
    if len(reader.pages) > 200:
        raise ValueError('PDF exceeds 200 pages')
    pages = []
    for number, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ''
        if len(text.strip()) < 40:
            output = path.parent / f'page-{number}'
            subprocess.run(
                [
                    'pdftoppm',
                    '-f',
                    str(number),
                    '-l',
                    str(number),
                    '-singlefile',
                    '-scale-to',
                    '2400',
                    '-png',
                    str(path),
                    str(output),
                ],
                check=True,
                capture_output=True,
                timeout=120,
            )
            text = ocr(output.with_suffix('.png'))
        pages.append(f'Page {number}\n{text}')
    return '\n\n'.join(pages)


def office_text(data: bytes, suffix: str) -> str:
    with ZipFile(io.BytesIO(data)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
            raise ValueError('Expanded document exceeds 100 MiB')
    if suffix == '.docx':
        from docx import Document

        doc = Document(io.BytesIO(data))
        return '\n'.join(
            [p.text for p in doc.paragraphs]
            + [' | '.join(c.text for c in row.cells) for table in doc.tables for row in table.rows]
        )
    if suffix == '.xlsx':
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            return '\n'.join(
                str(sheet.title)
                + '\n'
                + '\n'.join(
                    ' | '.join(str(c) if c is not None else '' for c in row)
                    for row in sheet.iter_rows(values_only=True)
                )
                for sheet in book
            )
        finally:
            book.close()
    from pptx import Presentation

    deck = Presentation(io.BytesIO(data))
    return '\n'.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)


def extract(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError('Unsupported document format')
    with tempfile.TemporaryDirectory(prefix='nawantara-extract-') as directory:
        path = Path(directory) / ('source' + suffix)
        path.write_bytes(data)
        if suffix == '.pdf':
            text = pdf_text(path)
        elif suffix in {'.png', '.jpg', '.jpeg', '.tif', '.tiff'}:
            text = ocr(path)
        elif suffix in {'.docx', '.xlsx', '.pptx'}:
            text = office_text(data, suffix)
        else:
            text = data.decode('utf-8-sig')
    if not text.strip():
        raise ValueError('No text could be extracted')
    if len(text) > 2_000_000:
        raise ValueError('Extracted text exceeds two million characters')
    return text


def chunks(text: str) -> list[str]:
    return [text[start : start + 1600] for start in range(0, len(text), 1400)]
